import sys
import subprocess
import os

# Auto-install python-docx if not installed
try:
    import docx
except ImportError:
    print("python-docx not found. Installing...")
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
    import docx

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_heading_styled(doc, text, level, space_before=12, space_after=6):
    heading = doc.add_heading(text, level=level)
    heading.paragraph_format.space_before = Pt(space_before)
    heading.paragraph_format.space_after = Pt(space_after)
    heading.paragraph_format.keep_with_next = True
    
    # Apply custom colors/sizes
    run = heading.runs[0]
    run.font.name = 'Rubik'
    if level == 1:
        run.font.size = Pt(20)
        run.font.bold = True
        run.font.color.rgb = RGBColor(2, 132, 199) # Primary #0284c7
    elif level == 2:
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = RGBColor(15, 23, 42)  # On Surface #0f172a
    elif level == 3:
        run.font.size = Pt(11.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(100, 116, 139) # Secondary #64748b
    return heading

def main():
    doc = Document()
    
    # Configure page margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        
    # Document Title Page
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(120)
    title_p.paragraph_format.space_after = Pt(10)
    
    run_title = title_p.add_run("ReAssure 3.0 Claims Auto-Adjudication Engine")
    run_title.font.name = 'Rubik'
    run_title.font.size = Pt(28)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(2, 132, 199)
    
    subtitle_p = doc.add_paragraph()
    subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_p.paragraph_format.space_after = Pt(250)
    run_sub = subtitle_p.add_run("Detailed System Architecture & Dynamic DAG Engine Specification")
    run_sub.font.name = 'Hanken Grotesk'
    run_sub.font.size = Pt(14)
    run_sub.font.color.rgb = RGBColor(100, 116, 139)
    
    meta_p = doc.add_paragraph()
    meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_meta = meta_p.add_run("Niva Bupa Policy Automation Team\nVersion 2.0.0\nConfidential - Internal Engineering Documentation")
    run_meta.font.name = 'Hanken Grotesk'
    run_meta.font.size = Pt(10)
    run_meta.font.italic = True
    run_meta.font.color.rgb = RGBColor(148, 163, 184)
    
    doc.add_page_break()
    
    # Page setup styling override for body
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Hanken Grotesk'
    style_normal.font.size = Pt(10.5)
    style_normal.font.color.rgb = RGBColor(15, 23, 42)
    style_normal.paragraph_format.line_spacing = 1.15
    style_normal.paragraph_format.space_after = Pt(6)
    
    # ── SECTION 1 ────────────────────────────────────────────────────────
    add_heading_styled(doc, "1. Executive Summary", 1)
    
    p = doc.add_paragraph(
        "The ReAssure 3.0 (R3) Claims Auto-Adjudication Engine is a high-throughput, "
        "low-latency system designed to automate the evaluation and payout of health insurance claims. "
        "Historically, claim adjudication required human auditors to manually cross-reference clinical "
        "discharge summaries against complex policy clause books, introducing delay and variance. "
        "This engine implements an AI-first paradigm: it pairs a deterministic 7-gate rule validation pipeline "
        "with an on-premise Large Language Model (LLM) for clinical semantic reasoning, running entirely within "
        "organizational network boundaries to satisfy compliance and data residency constraints."
    )
    
    p = doc.add_paragraph(
        "By dynamically building a Directed Acyclic Graph (DAG) of policy rules for every line item, "
        "topologically sorting the graph to enforce execution order, and parallelizing independent "
        "rule evaluations inside an asynchronous runtime, the system reduces typical claim decision latency "
        "to under 10 seconds. Human auditors are only engaged when composite rule evaluations fall below "
        "configurable confidence thresholds, establishing a secure four-tier routing framework."
    )
    
    # ── SECTION 2 ────────────────────────────────────────────────────────
    add_heading_styled(doc, "2. Subsystems & Decoupled Architecture", 1)
    
    p = doc.add_paragraph(
        "The application is structured into five decoupled processes to isolate execution contexts "
        "and scale computing workloads independently. These runtimes communicate via standard HTTP "
        "and server-sent events:"
    )
    
    # Table of processes
    table = doc.add_table(rows=6, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Light Shading Accent 1'
    
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Service / Process'
    hdr_cells[1].text = 'Default Port / Interface'
    hdr_cells[2].text = 'Core Responsibility'
    
    for i, cell in enumerate(hdr_cells):
        set_cell_background(cell, '0284C7')
        set_cell_margins(cell, top=120, bottom=120)
        for p_el in cell.paragraphs:
            for r in p_el.runs:
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
                
    processes_data = [
        ("PostgreSQL 16 Database", "5432 / TCP (asyncpg)", "Persistent storage for policies, members, balances, and historical claims."),
        ("llama.cpp Server", "8080 / HTTP (local REST)", "On-premise inference host running Gemma-4. Enforces strict JSON grammar schemas."),
        ("FastAPI Backend Engine", "8000 / HTTP & SSE", "Assembles contexts, computes DAG, executes rule validation gates, exposes API endpoints."),
        ("Vite Dev / Static Web Server", "5173 / HTTP (Web Interface)", "Hosts the single-page application dashboard for real-time visualization."),
        ("pgAdmin 4 Console", "5050 / HTTP (Web Admin)", "Optional web interface for direct database inspection and ledger state verification.")
    ]
    
    for idx, (svc, port, desc) in enumerate(processes_data):
        row = table.rows[idx + 1]
        row.cells[0].text = svc
        row.cells[1].text = port
        row.cells[2].text = desc
        bg_color = 'F8FAFC' if idx % 2 == 0 else 'FFFFFF'
        for cell in row.cells:
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=80, bottom=80)
            
    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    
    # ── SECTION 3 ────────────────────────────────────────────────────────
    add_heading_styled(doc, "3. Dynamic DAG Engine & The AI Planner", 1)
    
    p = doc.add_paragraph(
        "A critical design challenge in claim adjudication is that policy rules are highly interdependent. "
        "For example, you cannot calculate a room pro-rata deduction (Gate 6) until you have verified "
        "if the disease claimed is covered at all (Gate 3) and is not active in a waiting period (Gate 4). "
        "To solve this, the engine employs a dynamic Directed Acyclic Graph (DAG) compilation strategy."
    )
    
    add_heading_styled(doc, "3.1 Kahn's Algorithm & Topological Sort", 2)
    p = doc.add_paragraph(
        "For every line item, the AI Planner (src/planner.py) queries the Product Memory Store "
        "for rule blueprints applicable to the policy variant (Classic, Select, Elite). It dynamically "
        "builds a DAG by mapping the rule blueprints as nodes and their declared 'depends_on' arrays "
        "as directed edges. It then applies Kahn's algorithm for topological sorting:"
    )
    
    p = doc.add_paragraph(
        "1. Identify all nodes in the filtered set with an in-degree of 0 (no unresolved dependencies).\n"
        "2. Push these nodes into a Priority Queue. The priority key is a 3-tuple: (gate_value, rule_priority, rule_id), "
        "which enforces that rule processing respects the natural order of logical gates (Policy -> Member -> Coverage -> Waiting Period -> Exclusion -> Financial -> State).\n"
        "3. Systematically pop the node with the highest priority (lowest tuple value) from the queue, append it to the linear Execution Plan, "
        "and decrement the in-degree of its outgoing edges.\n"
        "4. If a target node's in-degree drops to 0, push it onto the Priority Queue.\n"
        "5. Repeat until the queue is empty. If the number of sorted nodes is less than the total node count, "
        "a dependency cycle is detected. The planner catches this safety violation, logs a warning, and falls back "
        "to a safe gate-sequential flat sort to ensure claim processing never stalls."
    )
    
    # ── SECTION 4 ────────────────────────────────────────────────────────
    add_heading_styled(doc, "4. Concurrency & Performance Model", 1)
    
    p = doc.add_paragraph(
        "Adjudication latency is minimized by executing independent steps concurrently. "
        "Because semantic rules require on-premise LLM inference (which takes 4 to 8 seconds depending on token counts "
        "and system hardware), synchronous sequential execution is unacceptable."
    )
    
    add_heading_styled(doc, "4.1 Depth Layer Grouping", 2)
    p = doc.add_paragraph(
        "After sorting the execution plan, the execution engine (src/pipeline_modules/execution.py) computes "
        "the dependency depth of each step in the graph. The depth is defined as the maximum path length from "
        "an in-degree 0 root node. Steps are then grouped into discrete depth layers:\n\n"
        "• Layer 0: Root checks that have 0 dependencies (e.g., Gate 1, Gate 2, initial deterministic validations).\n"
        "• Layer 1: Rules depending strictly on Layer 0 outputs.\n"
        "• Layer 2: Rules depending on Layer 1 outputs, and so on.\n\n"
        "Steps within the same depth layer are guaranteed to have zero dependency relationships with each other. "
        "Therefore, the engine fires all tasks in a layer concurrently using asyncio.gather()."
    )
    
    add_heading_styled(doc, "4.2 Thread Isolation & Non-blocking I/O", 2)
    p = doc.add_paragraph(
        "FastAPI runs on an asynchronous event loop (uvicorn). To prevent CPU-bound or blocking network calls "
        "from locking the event loop (which would cause the entire application to hang under concurrent load), "
        "the engine employs two distinct offloading models:\n\n"
        "1. Thread Pool Offloading: The SemanticExecutionAgent singleton dispatches all HTTP requests to "
        "the llama.cpp server inside standard Python threads via asyncio.run_in_executor(None, ...). "
        "This frees the main event loop to accept new incoming claim contexts while other threads await LLM answers.\n"
        "2. Stateful Sequential Processing: While rule execution happens concurrently, the aggregation of financial "
        "deductions (deductible consumption, room pro-rata, co-payment deductions) remains strictly sequential. "
        "The admissible and payable balances are updated one step at a time inside the zip(layer_steps, results) loop. "
        "This ensures that financial calculations are 100% deterministic and free of race conditions."
    )
    
    # ── SECTION 5 ────────────────────────────────────────────────────────
    add_heading_styled(doc, "5. The 7-Gate Validation Pipeline", 1)
    
    p = doc.add_paragraph(
        "Every claim context follows a strict sequence of 7 validation gates. "
        "If any gate evaluates to FAILED or EXCLUSION_ACTIVE, the pipeline halts immediately, "
        "skipping subsequent calculations to save processing cycles and returning the failure trace."
    )
    
    # Details of gates
    gates = [
        ("Gate 1: Policy Validation", "Deterministic", "Checks if the policy is active, the premium has been paid, if the claim falls within a grace period, and if the admission date falls between the policy start and end dates."),
        ("Gate 2: Member Validation", "Deterministic", "Verifies the member's status is active and that the admission date is equal to or greater than the date the member was added to the policy."),
        ("Gate 3: Coverage Validation", "Hybrid", "Executes policy-variant filters. Evaluates duration requirements (minimum 2 hours for daycare, 24 hours for alternative treatments). Triggers semantic LLM checks to verify if the treatment matches clinical coverage definitions."),
        ("Gate 4: Waiting Period Validation", "Hybrid", "Evaluates wait periods: 30-day initial wait (waived for accidents), 24-month specified-disease list, and 36-month Pre-Existing Disease (PED) wait. Subtracts porting credit months from active wait counters."),
        ("Gate 5: Exclusion Validation", "Semantic (LLM-driven)", "Checks the claim diagnosed condition and discharge summary against hard-coded list exclusions, then triggers the Semantic Agent to detect cosmetic treatments, diagnostic-only admissions, or out-patient exclusions."),
        ("Gate 6: Financial Computation", "Deterministic (Tools)", "Executes the 5-calculator waterfall sequence: applies room rent caps pro-rata, computes prolonged hospitalization co-pays, subtracts aggregate annual deductibles, and consumes SI pools (Base -> Booster+ -> ReAssure Forever)."),
        ("Gate 7: State Update", "Deterministic (Tools)", "Persists state updates: locks premium age under Lock the Clock rules, accumulates Booster+ credits for claim-free runs, computes Hospital Daily Cash benefits, and calculates Personal Accident payouts.")
    ]
    
    for title, mode, desc in gates:
        add_heading_styled(doc, title, 2)
        gp = doc.add_paragraph()
        r1 = gp.add_run(f"Execution Pattern: {mode}\n")
        r1.font.bold = True
        r1.font.color.rgb = RGBColor(100, 116, 139)
        gp.add_run(desc)
        
    doc.add_page_break()
    
    # ── SECTION 6 ────────────────────────────────────────────────────────
    add_heading_styled(doc, "6. Technical Diagram (Flowchart)", 1)
    
    p = doc.add_paragraph(
        "To render this diagram, copy the Mermaid syntax below and paste it into "
        "https://mermaid.live (the official online editor) or use a VS Code extension "
        "such as 'Markdown Preview Mermaid Support'."
    )
    
    # Add mermaid block in font-mono
    mermaid_block = doc.add_paragraph()
    mermaid_block.paragraph_format.left_indent = Inches(0.2)
    m_run = mermaid_block.add_run(
        "graph TD\n"
        "    Start([Claim Context Received]) --> G1[Gate 1: Policy Validation]\n"
        "    G1 -->|Passed| G2[Gate 2: Member Validation]\n"
        "    G1 -->|Failed| Exit[Early Exit / Rejected]\n"
        "    G2 -->|Passed| Plan[AI Planner: Filters & Builds DAG]\n"
        "    G2 -->|Failed| Exit\n"
        "    Plan --> Topological{Topological Sort}\n"
        "    Topological -->|Success| Exec[Asynchronous Layer-by-Layer Execution]\n"
        "    Topological -->|Cycle Detected| Fallback[Priority-Based Execution]\n"
        "    Exec --> G3[Gate 3: Coverage Validation]\n"
        "    Exec --> G4[Gate 4: Waiting Period]\n"
        "    Exec --> G5[Gate 5: Exclusion Validation]\n"
        "    G3 & G4 & G5 -->|All Passed| G6[Gate 6: Financial Calculators]\n"
        "    G3 & G4 & G5 -->|Any Exclusion/Fail| Exit\n"
        "    G6 --> G7[Gate 7: State Update & Persist]\n"
        "    G7 --> Dec([Generate Claim Decision])"
    )
    m_run.font.name = 'JetBrains Mono'
    m_run.font.size = Pt(9.5)
    m_run.font.color.rgb = RGBColor(9, 79, 114)
    
    # Save the file
    doc.save("claims_adjudication_technical_spec.docx")
    print("claims_adjudication_technical_spec.docx generated successfully.")

if __name__ == "__main__":
    main()
